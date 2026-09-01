#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
edgeR Analyzer (GUI) — A vs B (padronizado) — carbohydrate functional modules

Padronização
-----------
- edgeR logFC esperado: log2(A/B)
- Direção:
    logFC > 0  => A-up
    logFC < 0  => B-up
- Para gráficos divergentes por categoria:
    plotFC = -logFC  (A-up fica negativo; B-up fica positivo)

Requisitos
----------
  pip install pandas numpy matplotlib python-docx

Rodar
-----
  python3 edger_analyzer_gui_AB.py
"""

from __future__ import annotations

import re
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import Normalize
from matplotlib.cm import ScalarMappable

# ----------------------------
# Aparência dos gráficos
# ----------------------------
# Dimensões pensadas para figuras/pranchas de artigo:
# menores, mais retangulares e com texto proporcionalmente maior.
PLOT_DPI = 600

# Escala ÚNICA para todos os heatmaps z-score.
# Isso permite usar uma única barra de cores compartilhada na prancha.
HEATMAP_ZMIN = -2.0
HEATMAP_ZMAX = 2.0
HEATMAP_CMAP = "viridis"
HEATMAP_WIDTH = 5.2
BARPLOT_WIDTH = 5.4
TOPN_WIDTH = 5.4
CORR_WIDTH = 4.8
# Gráficos gerais de perfis médios por categoria
# - STANDALONE: inclui legenda externa, mas em canvas FIXO
# - PANEL: sem legenda, pensado para montagem de pranchas 2x2
PROFILE_STANDALONE_WIDTH = 7.2
PROFILE_PANEL_WIDTH = 5.8
PROFILE_HEIGHT = 4.2
MAX_LABEL_CHARS_COMPACT = 56

# Painéis fixos para manter todos os gráficos com o mesmo tamanho,
# independentemente do comprimento dos nomes dos genes.
CORR_HEIGHT = 4.6

# Escala Y fixa para comparação direta entre experimentos/fases
PROFILE_YMIN = -2.0
PROFILE_YMAX = 2.0
PROFILE_MARKER_SIZE = 5.2

# Altura física por gene/linha (em polegadas).
# A largura permanece fixa; a altura total cresce linearmente com n.
# Assim barras e células do heatmap mantêm SEMPRE a mesma altura.
HEATMAP_ROW_HEIGHT = 0.19
BAR_ROW_HEIGHT = 0.19
TOPN_ROW_HEIGHT = 0.19

# Margens verticais FIXAS em polegadas
HEATMAP_TOP_IN = 0.45
HEATMAP_BOTTOM_IN = 0.65
BAR_TOP_IN = 0.45
BAR_BOTTOM_IN = 0.65
TOPN_TOP_IN = 0.45
TOPN_BOTTOM_IN = 0.65
# Painel combinado: barplot + heatmap, gene a gene
# Largura fixa; altura física por gene também fixa.
COMBINED_WIDTH = 7.6
COMBINED_ROW_HEIGHT = 0.19
COMBINED_TOP_IN = 0.45
COMBINED_BOTTOM_IN = 0.70
COMBINED_LEFT = 0.36
COMBINED_RIGHT = 0.94
COMBINED_WSPACE = 0.08
# Margens fixas (evita que o tamanho final mude conforme o comprimento dos rótulos)
HEATMAP_LEFT = 0.50
HEATMAP_RIGHT = 0.87
HEATMAP_BOTTOM = 0.20
HEATMAP_TOP = 0.88

BAR_LEFT = 0.54
BAR_RIGHT = 0.95
BAR_BOTTOM = 0.17
BAR_TOP = 0.88

CORR_LEFT = 0.20
CORR_RIGHT = 0.90
CORR_BOTTOM = 0.20
CORR_TOP = 0.88

PROFILE_LEFT = 0.12
PROFILE_STANDALONE_RIGHT = 0.70
PROFILE_PANEL_RIGHT = 0.97
PROFILE_BOTTOM = 0.22
PROFILE_TOP = 0.88

# Fonte padrão dos gráficos:
# tenta Arial primeiro; se não existir no sistema, usa fallback compatível.
plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.sans-serif"] = ["Arial", "Liberation Sans", "DejaVu Sans"]

import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk


# ----------------------------
# Helpers: leitura e colunas
# ----------------------------

def _read_tabular(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, sep="\t", dtype=str, encoding="utf-8")
    except UnicodeDecodeError:
        return pd.read_csv(path, sep="\t", dtype=str, encoding="latin-1")


def _detect_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    cols_lower = {c.lower(): c for c in df.columns}
    for cand in candidates:
        if cand.lower() in cols_lower:
            return cols_lower[cand.lower()]
    return None


def _as_numeric_series(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s, errors="coerce")


def sanitize_basename(name: str) -> str:
    name = name.strip()
    if not name:
        return ""
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name)


# ----------------------------
# Filtro de descrições genéricas (opcional)
# ----------------------------

GENERIC_DESC_PATTERNS = [
    r"\buncharacterized\b",
    r"\bhypothetical\b",
    r"\bpredicted\b",
    r"\bunknown function\b",
    r"\bprotein of unknown function\b",
    r"\bunnamed protein product\b",
    r"\bputative uncharacterized\b",
    r"\bconserved unknown\b",
    r"\bno functional annotation\b",
]

def is_generic_description(desc: str) -> bool:
    if desc is None:
        return True
    d = str(desc).strip().lower()
    if d == "" or d == "nan":
        return True
    for p in GENERIC_DESC_PATTERNS:
        if re.search(p, d):
            return True
    if d.startswith("uncharacterized"):
        return True
    return False


def filter_informative_only(df: pd.DataFrame) -> pd.DataFrame:
    """Remove genes cuja descrição é genérica."""
    if "Description" not in df.columns:
        return df.copy()
    mask = ~df["Description"].apply(is_generic_description)
    return df.loc[mask].copy()


# ----------------------------
# Categorias (regex)
# ----------------------------

CATEGORY_RULES = [
    ("HSP_HSF_UPR_ER_Stress", [
        # Heat-shock proteins e resposta térmica
        r"heat[- ]shock protein",
        r"small heat[- ]shock protein",
        r"heat[- ]shock cognate",
        r"heat[- ]shock factor",
        r"heat[- ]stress transcription factor",
        r"heat shock \d+(?:\.\d+)? kDa protein",

        # HSPs nomeadas pelo símbolo/família
        r"\bhsp\d+(?:\.\d+)?[a-z0-9.-]*\b",
        r"\bhsc\d+(?:\.\d+)?[a-z0-9.-]*\b",
        r"\bshsp\b",

        # Principais famílias de chaperonas HSP
        r"\bhsp20\b",
        r"\bhsp40\b",
        r"\bhsp60\b",
        r"\bhsp70\b",
        r"\bhsp90\b",
        r"\bhsp100\b",
        r"\bhsp101\b",

        # DnaJ / HSP40 e J-domain proteins
        r"\bdnaj[a-z0-9.-]*\b",
        r"j[- ]domain protein",
        r"\berdj[a-z0-9.-]*\b",

        # Chaperoninas / desagregases / co-chaperonas
        r"\bchaperonin\b",
        r"\bcpn60[a-z0-9.-]*\b",
        r"\bclpb[a-z0-9.-]*\b",
        r"bag family molecular chaperone regulator",
        r"\bbag\d+[a-z0-9.-]*\b",
        r"\bmolecular chaperone\b",
        r"\bco[- ]?chaperone\b",

        # UPR / estresse do retículo endoplasmático
        r"unfolded protein response",
        r"\bUPR\b",
        r"endoplasmic reticulum stress",
        r"\bER stress\b",
        r"ER[- ]resident chaperone",
        r"ER[- ]luminal chaperone",

        # Sensores / fatores da UPR
        r"\bIRE1[A-Z0-9.-]*\b",
        r"\bbZIP60\b",
        r"\bbZIP28\b",

        # BiP / GRP78
        r"\bBiP\d*\b",
        r"\bGRP78\b",
        r"binding immunoglobulin protein",
        r"luminal binding protein",

        # Chaperonas / folding no ER
        r"calnexin",
        r"calreticulin",
        r"protein disulfide isomerase",
        r"\bPDI\d*\b",

        # ER-associated degradation
        r"ER[- ]associated degradation",
        r"\bERAD\b"
    ]),
    ("ROS_Antioxidant", [
        # Termos explícitos de ROS / estresse oxidativo
        r"oxidative stress",
        r"reactive oxygen species",
        r"\bROS\b",
        r"hydrogen peroxide",
        r"\bH2O2\b",

        # Superóxido dismutase
        r"superoxide dismutase",
        r"\bSOD\d*\b",
        r"Cu/Zn[- ]superoxide dismutase",
        r"Fe[- ]superoxide dismutase",
        r"Mn[- ]superoxide dismutase",

        # Catalase
        r"\bcatalase\b",
        r"\bCAT\d*\b",

        # Peroxidases
        r"ascorbate peroxidase",
        r"\bAPX\d*\b",
        r"glutathione peroxidase",
        r"\bGPX\d*\b",
        r"peroxidase",

        # Ciclo ascorbato-glutationa
        r"monodehydroascorbate reductase",
        r"\bMDHAR\d*\b",
        r"dehydroascorbate reductase",
        r"\bDHAR\d*\b",
        r"glutathione reductase",
        r"ascorbate[- ]glutathione",

        # Glutationa / detoxificação redox
        r"glutathione S[- ]transferase",
        r"glutathione transferase",
        r"\bGST[A-Z0-9.-]*\b",
        r"glutathione",

        # Sistemas tiol-redox
        r"thioredoxin",
        r"\bTRX[A-Z0-9.-]*\b",
        r"glutaredoxin",
        r"\bGRX[A-Z0-9.-]*\b",
        r"peroxiredoxin",
        r"\bPRX[A-Z0-9.-]*\b",

        # Outros sistemas associados a ROS
        r"alternative oxidase",
        r"\bAOX\d*\b"
    ]),
    ("Sucrose_Synthesis_Maintenance", [
        # -------------------------
        # Sucrose synthesis / maintenance
        # -------------------------
        # SPS + SPP form the canonical sucrose-phosphate synthesis route.
        r"sucrose[- ]phosphate synthase",
        r"\bSPS\b",
        r"\bSPS\d+[A-Za-z0-9.-]*\b",
        r"sucrose[- ]phosphate phosphatase",
        r"\bSPP\b",
        r"\bSPP\d+[A-Za-z0-9.-]*\b",

        # Invertase inhibitors are retained here because they can reduce
        # sucrose hydrolysis and therefore favor sucrose maintenance.
        # Only explicit inhibitor annotations are included.
        r"\binvertase[- ]inhibitor\b",
        r"\binvertase inhibitor\b",
    ]),
    ("Sucrose_Cleavage_Utilization", [
        # -------------------------
        # Sucrose cleavage / utilization
        # -------------------------
        # Sucrose synthase is reversible. In sink/reproductive tissues it is
        # commonly associated with sucrose cleavage/utilization, but individual
        # genes should still be interpreted with that biochemical caveat.
        r"sucrose[- ]synthase",
        r"\bSUS\b",
        r"\bSUS\d+[A-Za-z0-9.-]*\b",
        r"\bSuSy\b",
        r"\bSuSy\d+[A-Za-z0-9.-]*\b",

        # Invertases / beta-fructofuranosidases
        r"\binvertase\b",
        r"beta[- ]fructofuranosidase",
        r"cell wall invertase",
        r"\bCWINV\b",
        r"\bCWINV\d+[A-Za-z0-9.-]*\b",
        r"vacuolar invertase",
        r"\bVIN\b",
        r"\bVIN\d+[A-Za-z0-9.-]*\b",
        r"neutral invertase",
        r"cytosolic invertase",
        r"\bCINV\b",
        r"\bCINV\d+[A-Za-z0-9.-]*\b",
        r"\bNINV\b",
        r"\bNINV\d+[A-Za-z0-9.-]*\b",
    ]),
    ("Starch_Synthesis_Maintenance", [
        # -------------------------
        # Starch synthesis / maintenance
        # -------------------------
        r"starch synthase",
        r"granule[- ]bound starch synthase",
        r"\bGBSS\b",
        r"\bGBSS\d+[A-Za-z0-9.-]*\b",
        r"ADP[- ]glucose pyrophosphorylase",
        r"ADP glucose pyrophosphorylase",
        r"\bAGPase\b",
        r"glucose[- ]1[- ]phosphate adenylyltransferase",
        r"starch branching enzyme",
        r"\bSBE\b",
        r"\bSBE\d+[A-Za-z0-9.-]*\b",

        # A true alpha-amylase inhibitor is a direct negative regulator of
        # starch degradation and is therefore grouped with starch maintenance.
        r"\balpha[- ]amylase inhibitor\b",
        r"\balpha amylase inhibitor\b",
    ]),
    ("Starch_Degradation_Mobilization", [
        # -------------------------
        # Starch remodeling / debranching
        # -------------------------
        r"isoamylase",
        r"\bISA\b",
        r"\bISA\d+[A-Za-z0-9.-]*\b",
        r"pullulanase",
        r"limit dextrinase",

        # -------------------------
        # Starch degradation / mobilization
        # -------------------------
        r"alpha[- ]amylase",
        r"beta[- ]amylase",
        r"\bamylase\b",
        r"\bAMY\b",
        r"\bAMY\d+[A-Za-z0-9.-]*\b",
        r"\bBAM\b",
        r"\bBAM\d+[A-Za-z0-9.-]*\b",
        r"glucan water dikinase",
        r"\bGWD\b",
        r"\bGWD\d+[A-Za-z0-9.-]*\b",
        r"phosphoglucan water dikinase",
        r"\bPWD\b",
        r"\bPWD\d+[A-Za-z0-9.-]*\b",
        r"starch excess",
        r"\bSEX1\b",
        r"\bSEX4\b",
        r"disproportionating enzyme",
        r"\bDPE\b",
        r"\bDPE\d+[A-Za-z0-9.-]*\b",
        r"glucan phosphorylase",
        r"starch phosphorylase",
        r"alpha[- ]glucan phosphorylase",
        r"\bPHS\b",
        r"\bPHS\d+[A-Za-z0-9.-]*\b",
    ]),
    ("Sugar_Transport", [
        # -------------------------
        # Sugar transport
        # -------------------------
        # SUT/SUC are canonical sucrose transporters.
        r"sucrose transporter",
        r"sucrose carrier",
        r"\bSUT\b",
        r"\bSUT\d+[A-Za-z0-9.-]*\b",
        r"\bSUC\b",
        r"\bSUC\d+[A-Za-z0-9.-]*\b",

        # SWEETs are kept in a separate transport module because substrate
        # preference varies among family members (sucrose vs hexoses, etc.).
        # Their physiological direction also depends on tissue/localization
        # and concentration gradients, so no accumulation/depletion direction
        # is inferred automatically from expression alone.
        r"\bSWEET\b",
        r"\bSWEET\d+[A-Za-z0-9.-]*\b",
    ]),
    ("Metabolism", [
        r"glycolysis", r"gluconeogenesis", r"tricarboxylic", r"\bTCA\b", r"citric acid cycle",
        r"respiration", r"mitochondrial", r"oxidative phosphorylation", r"\bATP synthase\b",
        r"malate dehydrogenase", r"isocitrate dehydrogenase", r"succinate dehydrogenase",
        r"pyruvate kinase", r"hexokinase", r"phosphofructokinase",
        r"shikimate", r"phenylpropanoid", r"flavonoid", r"terpen", r"carotenoid",
        r"fatty acid", r"lipid", r"beta-oxidation"
    ]),
    ("Ubiquitin_Proteasome", [
        # -------------------------
        # Ubiquitina / poliubiquitina
        # -------------------------
        r"\bubiquitin\b",
        r"\bpolyubiquitin\b",
        r"ubiquitin[- ]like protein",

        # -------------------------
        # E1 — ativação da ubiquitina
        # -------------------------
        r"ubiquitin[- ]activating enzyme",
        r"\bE1 ubiquitin",
        r"\bUBA\d*\b",

        # -------------------------
        # E2 — conjugação
        # -------------------------
        r"ubiquitin[- ]conjugating enzyme",
        r"\bE2 ubiquitin",
        r"\bUBC\d*\b",

        # -------------------------
        # E3 — ligases
        # -------------------------
        r"E3 ubiquitin[- ]protein ligase",
        r"E3 ubiquitin ligase",
        r"ubiquitin[- ]protein ligase",
        r"ubiquitin ligase",

        # RING E3
        r"\bRING[- ]finger\b",
        r"\bRING[- ]H2\b",
        r"\bRING[- ]HC\b",
        r"\bRING[- ]type\b",
        r"\bRING E3\b",

        # U-box / PUB
        r"\bU[- ]box\b",
        r"U[- ]box domain",
        r"\bPUB\d*\b",

        # -------------------------
        # SCF / F-box
        # -------------------------
        r"\bF[- ]box\b",
        r"F[- ]box protein",
        r"F[- ]box/kelch",
        r"\bSCF complex\b",
        r"SKP1[- ]like",
        r"\bSKP1\b",

        # -------------------------
        # CUL3 / BTB
        # -------------------------
        r"BTB/POZ",
        r"BTB[- ]domain",
        r"BTB domain",
        r"CUL3[- ]BTB",
        r"\bCullin[- ]3\b",

        # -------------------------
        # Cullin-RING ligases
        # -------------------------
        r"\bcullin\b",
        r"Cullin[- ]RING",
        r"\bRBX1\b",

        # -------------------------
        # Proteassoma
        # -------------------------
        r"\bproteasome\b",
        r"26S proteasome",
        r"20S proteasome",
        r"proteasome regulatory",
        r"proteasome subunit",

        # -------------------------
        # Desubiquitinação
        # -------------------------
        r"deubiquitin",
        r"deubiquitylat",
        r"ubiquitin carboxyl[- ]terminal hydrolase",
        r"ubiquitin[- ]specific protease",
        r"\bUBP\d*\b"
    ]),
    ("Cytoskeleton_Polarity_Trafficking", [
        # -------------------------
        # Microtúbulos
        # -------------------------
        r"microtubul",
        r"\btubulin\b",
        r"\bTUA\d*\b",
        r"\bTUB\d*\b",
        r"kinesin",
        r"dynein",
        r"\bMAP65\b",
        r"katanin",
        r"microtubule[- ]associated protein",
        r"microtubule[- ]destabilizing protein",
        r"EB family member",

        # -------------------------
        # Actina e proteínas associadas
        # -------------------------
        r"\bactin-\d+\b",
        r"\bACT\d+\b",
        r"actin cytoskeleton",
        r"actin filament",
        r"actin[- ]binding protein",
        r"actin-related protein",
        r"actin depolymerizing factor",
        r"profilin",
        r"cofilin",

        # IMPORTANTE: \b impede capturar "UDP-forming"
        r"\bformin(?:[- ]like)?\b",

        r"fimbrin",
        r"villin",
        r"myosin",

        # -------------------------
        # Polaridade celular
        # -------------------------
        r"cell polarity",
        r"polarized growth",
        r"polar growth",
        r"polar localization",
        r"ROP GTPase",
        r"Rho[- ]related GTPase",
        r"Rho of plants",
        r"RAC/ROP",
        r"SOSEKI",
        r"DUF966",

        # -------------------------
        # Pequenas GTPases de tráfego
        # -------------------------
        r"Rab GTPase",
        r"Rab-like GTPase",
        r"Rab family",
        r"Rab GTPase[- ]activating protein",
        r"Rab[- ]GAP",

        # ARF somente quando funcionalmente explícito
        r"ADP[- ]ribosylation factor",
        r"ARF GTPase",
        r"ARF[- ]GEF",
        r"ARF[- ]GAP",

        # -------------------------
        # Tráfego vesicular
        # -------------------------
        r"\bSNARE\b",
        r"syntaxin",
        r"synaptobrevin",
        r"\bVAMP\d*\b",
        r"exocyst",
        r"\bEXO70[A-Z0-9.-]*\b",
        r"\bEXO84[A-Z0-9.-]*\b",
        r"coatomer",
        r"\bCOPI\b",
        r"\bCOPII\b",
        r"clathrin",
        r"dynamin",
        r"endocyt",
        r"vesicle trafficking",
        r"vesicular trafficking",
        r"vesicle transport",
        r"vesicular transport",
        r"membrane trafficking",
        r"Golgi trafficking"
    ]),
    ("CellWall_Callose_Pectin", [
        # Termos gerais
        r"cell wall",
        r"cell-wall",
        r"wall remodeling",
        r"wall modification",

        # Calose / glucanos
        r"callose",
        r"callose synthase",
        r"glucan synthase",
        r"glucan synthase-like",
        r"\bGSL\d*\b",
        r"beta-1,3-glucan",
        r"β-1,3-glucan",

        # Celulose
        r"cellulose",
        r"cellulose synthase",
        r"\bCESA\d*\b",

        # Xiloglucano
        r"xyloglucan",
        r"xyloglucan endotransglucosylase",
        r"xyloglucan endotransglycosylase",
        r"xyloglucan hydrolase",
        r"\bXTH\d*\b",

        # Pectina
        r"pectin",
        r"pectinesterase",
        r"pectin esterase",
        r"pectin methylesterase",
        r"pectin methyl esterase",
        r"\bPME\d*\b",
        r"polygalacturonase",
        r"pectate lyase",
        r"pectin lyase",

        # Arabinano / arabinose
        r"arabinan",
        r"arabinofuranosidase",
        r"alpha-L-arabinofuranosidase",
        r"α-L-arabinofuranosidase",

        # Xilano / xilose
        r"\bxylan\b",
        r"xylanase",
        r"xylosidase",
        r"beta-xylosidase",
        r"β-xylosidase",

        # Outros polissacarídeos da parede
        r"mannan",
        r"mannanase",
        r"galactan",
        r"galactanase",
        r"arabinogalactan",

        # Proteínas estruturais/remodeladoras
        r"expansin",
        r"extensin",

        # Parede secundária
        r"lignin",
        r"lignification"
    ]),
    ("Hormone_Development", [
        r"auxin", r"\bIAA\b", r"gibberellin", r"\bGA\b", r"cytokinin",
        r"abscisic", r"\bABA\b", r"ethylene", r"\bACC\b", r"jasmon", r"salicylic",
        r"brassin", r"floral", r"pollen", r"anther", r"tapetum", r"microspore", r"DELLA"
    ]),
    ("TF_Chromatin", [
        r"\bMYB\b", r"\bbHLH\b", r"\bNAC\b", r"\bWRKY\b", r"\bMADS\b", r"\bbZIP\b",
        r"chromatin", r"histone", r"SWI/SNF", r"methyltransferase", r"acetyltransferase", r"deacetylase"
    ]),
]

CATEGORY_NAMES = [cat for cat, _ in CATEGORY_RULES]

# Categorias disponíveis para seleção na GUI.
# "Other" não pertence a CATEGORY_RULES porque é a categoria de fallback,
# mas pode ser selecionada manualmente quando o usuário quiser gerar seus gráficos.
SELECTABLE_CATEGORIES = CATEGORY_NAMES + ["Other"]

# Ordem/cores FIXAS para os gráficos gerais. Assim uma categoria mantém
# a mesma cor mesmo quando outra categoria está ausente em uma comparação.
PROFILE_CATEGORY_ORDER = [
    "HSP_HSF_UPR_ER_Stress",
    "ROS_Antioxidant",
    "Ubiquitin_Proteasome",
    "Cytoskeleton_Polarity_Trafficking",
    "CellWall_Callose_Pectin",
    "Sucrose_Synthesis_Maintenance",
    "Sucrose_Cleavage_Utilization",
    "Starch_Synthesis_Maintenance",
    "Starch_Degradation_Mobilization",
    "Sugar_Transport",
    "Metabolism",
    "Hormone_Development",
    "TF_Chromatin",
    "Other",
]
_PROFILE_CMAP = plt.get_cmap("tab10")
_PROFILE_CMAP20 = plt.get_cmap("tab20")
# Preserva exatamente as cores das categorias que já existiam na v7.
# Os cinco módulos de carboidratos recebem cores próprias sem deslocar as antigas.
PROFILE_CATEGORY_COLORS = {
    "HSP_HSF_UPR_ER_Stress": _PROFILE_CMAP(0),
    "ROS_Antioxidant": _PROFILE_CMAP(1),
    "Ubiquitin_Proteasome": _PROFILE_CMAP(2),
    "Cytoskeleton_Polarity_Trafficking": _PROFILE_CMAP(3),
    "CellWall_Callose_Pectin": _PROFILE_CMAP(4),
    "Metabolism": _PROFILE_CMAP(5),
    "Hormone_Development": _PROFILE_CMAP(6),
    "TF_Chromatin": _PROFILE_CMAP(7),
    "Other": _PROFILE_CMAP(8),
    "Sucrose_Synthesis_Maintenance": _PROFILE_CMAP20(10),
    "Sucrose_Cleavage_Utilization": _PROFILE_CMAP20(11),
    "Starch_Synthesis_Maintenance": _PROFILE_CMAP20(12),
    "Starch_Degradation_Mobilization": _PROFILE_CMAP20(13),
    "Sugar_Transport": _PROFILE_CMAP20(14),
}
PROFILE_CATEGORY_LABELS = {
    "HSP_HSF_UPR_ER_Stress": "HSP / UPR",
    "ROS_Antioxidant": "ROS / antioxidant",
    "Ubiquitin_Proteasome": "Ubiquitin / proteasome",
    "Cytoskeleton_Polarity_Trafficking": "Cytoskeleton / trafficking",
    "CellWall_Callose_Pectin": "Cell wall",
    "Sucrose_Synthesis_Maintenance": "Sucrose synthesis / maintenance",
    "Sucrose_Cleavage_Utilization": "Sucrose cleavage / utilization",
    "Starch_Synthesis_Maintenance": "Starch synthesis / maintenance",
    "Starch_Degradation_Mobilization": "Starch degradation / mobilization",
    "Sugar_Transport": "Sugar transport",
    "Metabolism": "Metabolism",
    "Hormone_Development": "Hormone / development",
    "TF_Chromatin": "TF / chromatin",
    "Other": "Other",
}


# Falsos positivos confirmados por auditoria manual.
# A exclusão é ESPECÍFICA DA CATEGORIA: o gene não é removido do dataset e ainda
# pode ser classificado em outra categoria caso satisfaça legitimamente suas regras.
CATEGORY_EXCLUDED_GENEIDS = {
    "HSP_HSF_UPR_ER_Stress": {"101245517"},
    "ROS_Antioxidant": {"101267166", "101267453"},
    "Cytoskeleton_Polarity_Trafficking": {"101266703"},
    "CellWall_Callose_Pectin": {"101255470"},
    # O texto legado "ethylene-inducing xylanase" também acionaria indevidamente
    # a regra ampla de etileno em Hormone_Development após a exclusão de CellWall.
    "Hormone_Development": {"101255470"},

    # LOC101245586 = alpha-amylase inhibitor/lipid-transfer/seed-storage
    # superfamily protein. O nome da SUPERFAMÍLIA não demonstra que esta proteína
    # específica seja um inibidor funcional de alpha-amylase. Portanto, ela não
    # deve entrar automaticamente em starch synthesis/maintenance.
    "Starch_Synthesis_Maintenance": {"101245586"},
    "Starch_Degradation_Mobilization": {"101245586"},
    "Metabolism": {"101245586"},
}


# Padrões negativos curados para separar enzimas de seus inibidores e evitar
# falsos positivos por nomes de superfamílias.
CATEGORY_EXCLUDE_PATTERNS = {
    # Um invertase inhibitor pertence a sucrose synthesis/maintenance, não ao
    # módulo de cleavage/utilization só porque contém a palavra "invertase".
    "Sucrose_Cleavage_Utilization": [
        r"invertase[- ]inhibitor",
        r"invertase inhibitor",
    ],

    # Um alpha-amylase inhibitor pertence a starch synthesis/maintenance, não
    # ao módulo de degradação só porque contém "alpha-amylase".
    "Starch_Degradation_Mobilization": [
        r"alpha[- ]amylase inhibitor",
        r"alpha amylase inhibitor",
        r"amylase[- ]inhibitor",
        r"amylase inhibitor",
    ],

    # Porém descrições que indicam apenas a superfamília estrutural
    # alpha-amylase inhibitor/lipid-transfer/seed-storage NÃO são suficientes
    # para assumir função de inibidor de amilase.
    "Starch_Synthesis_Maintenance": [
        r"amylase inhibitor/lipid[- ]transfer/seed storage",
        r"amylase inhibitor/lipid[- ]transfer/seed[- ]storage",
    ],

    # O mesmo falso positivo contém "lipid-transfer" e poderia cair na categoria
    # ampla Metabolism apenas pela palavra "lipid".
    "Metabolism": [
        r"amylase inhibitor/lipid[- ]transfer/seed storage",
        r"amylase inhibitor/lipid[- ]transfer/seed[- ]storage",
    ],
}


def categorize_row(geneid: str, symbol: str, desc: str) -> list[str]:
    gid = str(geneid or "").strip()
    s = ((symbol or "") + " " + (desc or "")).lower()
    cats = []
    for cat, patterns in CATEGORY_RULES:
        # Não permitir falsos positivos auditados nesta categoria específica.
        if gid in CATEGORY_EXCLUDED_GENEIDS.get(cat, set()):
            continue

        # Bloquear descrições que só contêm o nome da enzima em contexto de
        # INIBIDOR/superfamília, evitando falsos positivos lexicais.
        if any(re.search(p.lower(), s) for p in CATEGORY_EXCLUDE_PATTERNS.get(cat, [])):
            continue

        for p in patterns:
            if re.search(p.lower(), s):
                cats.append(cat)
                break
    if not cats:
        cats = ["Other"]
    return cats


# ----------------------------
# DETECÇÃO ROBUSTA DAS AMOSTRAS (3, 4 ou 5; A e B com mesmo n)
# ----------------------------

def detect_and_normalize_sample_cols(counts_df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """
    Mantém exatamente o comportamento para 3 replicatas e expande para 4/5 quando existir.
    Regras:
      - aceita n = 3, 4 ou 5
      - A e B DEVEM ter o mesmo n
      - aceita colunas A1..An,B1..Bn OU Counts_A1..Counts_Bn (case-insensitive)
    Retorna:
      - df com colunas renomeadas para A1..An,B1..Bn
      - sample_cols na ordem A1..An,B1..Bn
    """
    cols = list(counts_df.columns)

    def wanted_for(n: int) -> list[str]:
        return [f"A{i}" for i in range(1, n + 1)] + [f"B{i}" for i in range(1, n + 1)]

    # 1) Caso já esteja normalizado (A1.. / B1..) — prioriza 5, depois 4, depois 3
    for n in (5, 4, 3):
        wanted = wanted_for(n)
        if all(c in cols for c in wanted):
            return counts_df.copy(), wanted

    # 2) Caso esteja com prefixo Counts_ (Counts_A1..Counts_Bn)
    for n in (5, 4, 3):
        wanted = wanted_for(n)
        mapping = {}
        for w in wanted:
            for pref in ("Counts_", "counts_"):
                key = f"{pref}{w}"
                if key in cols:
                    mapping[key] = w
        if len(mapping) == len(wanted):
            df = counts_df.copy().rename(columns=mapping)
            return df, wanted

    # 3) Caso misto / variações de caixa: regex ^(?:Counts_)?([AB][1-5])$
    regex = re.compile(r"^(?:Counts_)?([AB][1-5])$", flags=re.IGNORECASE)
    found = {}
    for c in cols:
        m = regex.match(c)
        if m:
            found[c] = m.group(1).upper()

    # escolher o maior n em {5,4,3} que tenha conjunto completo e pareado
    for n in (5, 4, 3):
        wanted = set(wanted_for(n))
        if set(found.values()) == wanted:
            df = counts_df.copy().rename(columns=found)
            return df, wanted_for(n)

    raise ValueError(
        "Counts: não consegui identificar colunas para A vs B com n=3,4 ou 5 replicatas.\n"
        f"Colunas encontradas: {cols}\n\n"
        "Garanta colunas exatamente:\n"
        "  - 'A1 A2 A3 B1 B2 B3' (ou até A5/B5),\n"
        "OU\n"
        "  - 'Counts_A1 ... Counts_B3' (ou até Counts_A5/Counts_B5).\n"
        "E lembre: A e B devem ter o mesmo número de replicatas (3, 4 ou 5)."
    )


# ----------------------------
# CPM / logCPM / z-score
# ----------------------------

def compute_log2cpm_from_counts(counts_df: pd.DataFrame, sample_cols: list[str]) -> pd.DataFrame:
    X = counts_df[sample_cols].apply(pd.to_numeric, errors="coerce").fillna(0.0).astype(float)
    libsize = X.sum(axis=0).replace(0, np.nan)
    cpm = (X / libsize) * 1e6
    log2cpm = np.log2(cpm + 1.0)
    out = counts_df[["GeneID"]].copy()
    out[sample_cols] = log2cpm
    return out


def zscore_rows(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    X = df[cols].astype(float).values
    mu = np.nanmean(X, axis=1, keepdims=True)
    sd = np.nanstd(X, axis=1, keepdims=True)
    sd[sd == 0] = 1.0
    Z = (X - mu) / sd
    out = df.copy()
    out[cols] = Z
    return out


# ----------------------------
# Labels
# ----------------------------

def _build_desc_labels(desc_series: pd.Series, symbol_series: pd.Series | None = None,
                      max_chars: int = 140) -> pd.Series:
    """
    Rótulos por DESCRIÇÃO, com fallback para símbolo.
    """
    dsc = desc_series.fillna("").astype(str).replace("nan", "").str.strip()
    sym = symbol_series.fillna("").astype(str).replace("nan", "").str.strip() if symbol_series is not None else pd.Series([""] * len(dsc), index=dsc.index)
    labels = dsc.copy()
    empty = labels.eq("") | labels.str.lower().eq("nan")
    labels.loc[empty] = sym.loc[empty]
    labels = labels.replace("", "NA")
    labels = labels.apply(lambda x: x if len(x) <= max_chars else x[: max_chars - 3] + "...")
    return labels


def _build_symbol_labels(symbol_series: pd.Series,
                         geneid_series: pd.Series | None = None,
                         max_chars: int = 140) -> pd.Series:
    """
    Rótulos por SÍMBOLO, com fallback para GeneID.
    """
    sym = symbol_series.fillna("").astype(str).replace("nan", "").str.strip()
    if geneid_series is not None:
        gid = geneid_series.fillna("").astype(str).replace("nan", "").str.strip()
    else:
        gid = pd.Series([""] * len(sym), index=sym.index)

    labels = sym.copy()
    empty = labels.eq("") | labels.str.lower().eq("nan")
    labels.loc[empty] = gid.loc[empty]
    labels = labels.replace("", "NA")
    labels = labels.apply(lambda x: x if len(x) <= max_chars else x[: max_chars - 3] + "...")
    return labels


# ----------------------------
# PLOTS — HEATMAPS
# ----------------------------

def save_heatmap(
    matrix_df: pd.DataFrame,
    out_png: Path,
    title: str,
    sample_cols: list[str],
    row_labels: pd.Series,
    max_rows: int = 300,
    also_pdf: bool = True,
):
    # NÃO MEXER: usa SOMENTE sample_cols
    X = matrix_df[sample_cols].astype(float).values

    if X.shape[0] > max_rows:
        var = np.nanvar(X, axis=1)
        idx = np.argsort(var)[::-1][:max_rows]
        X = X[idx, :]
        row_labels = row_labels.iloc[idx]

    n_rows = X.shape[0]
    n_cols = len(sample_cols)

    # Largura fixa e estreita; altura cresce com o número de genes
    fig_h = HEATMAP_TOP_IN + HEATMAP_BOTTOM_IN + HEATMAP_ROW_HEIGHT * n_rows
    fig, ax = plt.subplots(figsize=(HEATMAP_WIDTH, fig_h))
    im = ax.imshow(
        X,
        aspect="auto",
        origin="upper",
        interpolation="nearest",
        cmap=HEATMAP_CMAP,
        vmin=HEATMAP_ZMIN,
        vmax=HEATMAP_ZMAX,
    )

    cbar = fig.colorbar(im, ax=ax, fraction=0.045, pad=0.02)
    cbar.set_label("z-score", fontsize=8.5)
    cbar.ax.tick_params(labelsize=7.5)

    ax.set_xticks(np.arange(n_cols))
    ax.set_xticklabels(sample_cols, rotation=45, ha="right", fontsize=8)

    labels = row_labels.astype(str).values

    # Mostrar todos os genes e manter a mesma fonte;
    # a altura do gráfico é que cresce quando há mais genes.
    tick_idx = np.arange(n_rows)
    y_fs = 7.0

    ax.set_yticks(tick_idx)
    ax.set_yticklabels(labels[tick_idx], fontsize=y_fs)

    ax.set_title(title, fontsize=9.3, pad=7)
    ax.tick_params(axis="both", length=3)

    fig.subplots_adjust(
        left=HEATMAP_LEFT, right=HEATMAP_RIGHT,
        bottom=HEATMAP_BOTTOM_IN / fig_h, top=1.0 - (HEATMAP_TOP_IN / fig_h)
    )

    fig.savefig(out_png, dpi=PLOT_DPI)
    if also_pdf:
        fig.savefig(out_png.with_suffix(".pdf"))
    plt.close(fig)


def save_combined_bar_heatmap(
    df_cat: pd.DataFrame,
    out_png: Path,
    title: str,
    sample_cols: list[str],
    label_mode: str = "desc",
    max_label_chars: int = MAX_LABEL_CHARS_COMPACT,
):
    """
    Painel combinado, gene a gene:
      ESQUERDA: barplot divergente (plotFC = -logFC)
      DIREITA: heatmap de z-score nas amostras

    A mesma linha representa exatamente o mesmo gene nos dois painéis.
    Ordem visual:
      topo   = B-up mais forte
      base   = A-up mais forte
    """
    if df_cat.empty:
        return

    d = df_cat.copy()
    d["logFC"] = pd.to_numeric(d["logFC"], errors="coerce")
    d = d.dropna(subset=["logFC"]).copy()
    if d.empty:
        return

    # Mesma convenção já usada nos barplots
    d["plotFC"] = -d["logFC"]

    # Para o heatmap, a primeira linha aparece no topo.
    # logFC crescente => B-up mais forte no topo e A-up mais forte embaixo.
    d = d.sort_values("logFC", ascending=True).reset_index(drop=True)

    if label_mode == "symbol":
        labels = _build_symbol_labels(
            d["Symbol"], d.get("GeneID"), max_chars=max_label_chars
        )
    else:
        labels = _build_desc_labels(
            d["Description"], d.get("Symbol"), max_chars=max_label_chars
        )

    X = d[sample_cols].astype(float).values
    n = len(d)
    n_cols = len(sample_cols)

    # Mesma escala vertical por gene usada nos barplots.
    fig_h = COMBINED_TOP_IN + COMBINED_BOTTOM_IN + COMBINED_ROW_HEIGHT * n
    fig = plt.figure(figsize=(COMBINED_WIDTH, fig_h))

    gs = fig.add_gridspec(
        1, 2,
        width_ratios=[1.0, 1.28],
        left=COMBINED_LEFT,
        right=COMBINED_RIGHT,
        bottom=COMBINED_BOTTOM_IN / fig_h,
        top=1.0 - (COMBINED_TOP_IN / fig_h),
        wspace=COMBINED_WSPACE,
    )

    ax_bar = fig.add_subplot(gs[0, 0])
    ax_hm = fig.add_subplot(gs[0, 1], sharey=ax_bar)

    # ----------------
    # Barplot
    # ----------------
    y = np.arange(n)
    ax_bar.barh(
        y,
        d["plotFC"].astype(float).values,
        height=0.72,
    )
    ax_bar.axvline(0, linewidth=0.9)

    ax_bar.set_yticks(y)
    ax_bar.set_yticklabels(labels.astype(str).values, fontsize=7.0)
    ax_bar.set_xlabel(
        "plotFC = -log2FC(A/B)\n(A-up ← | → B-up)",
        fontsize=8.2,
    )
    ax_bar.tick_params(axis="x", labelsize=7.6)
    ax_bar.tick_params(axis="y", labelsize=7.0)

    # Linha 0 deve ficar no topo, igual ao imshow(origin='upper')
    ax_bar.set_ylim(n - 0.5, -0.5)

    # ----------------
    # Heatmap
    # ----------------
    im = ax_hm.imshow(
        X,
        aspect="auto",
        origin="upper",
        interpolation="nearest",
        cmap=HEATMAP_CMAP,
        vmin=HEATMAP_ZMIN,
        vmax=HEATMAP_ZMAX,
    )

    ax_hm.set_xticks(np.arange(n_cols))
    ax_hm.set_xticklabels(
        sample_cols,
        rotation=45,
        ha="right",
        fontsize=7.8,
    )

    # Não repetir os nomes dos genes no lado do heatmap
    ax_hm.tick_params(
        axis="y",
        which="both",
        left=False,
        labelleft=False,
    )
    ax_hm.set_ylim(n - 0.5, -0.5)

    cbar = fig.colorbar(
        im,
        ax=ax_hm,
        fraction=0.055,
        pad=0.025,
    )
    cbar.set_label("z-score", fontsize=8.2)
    cbar.ax.tick_params(labelsize=7.4)

    fig.suptitle(title, fontsize=9.4, y=0.975)

    fig.savefig(out_png, dpi=PLOT_DPI)
    fig.savefig(out_png.with_suffix(".pdf"))
    plt.close(fig)


def save_shared_zscore_scales(out_dir: Path, base: str):
    """
    Gera barras de cores independentes usando EXATAMENTE a mesma escala
    e o mesmo cmap de todos os heatmaps. Produz versão vertical e horizontal.

    A versão horizontal é invertida (valores positivos à esquerda) para ficar
    visualmente equivalente à prancha montada pelo usuário.
    """
    norm = Normalize(vmin=HEATMAP_ZMIN, vmax=HEATMAP_ZMAX)
    sm = ScalarMappable(norm=norm, cmap=HEATMAP_CMAP)
    sm.set_array([])
    ticks = np.arange(HEATMAP_ZMIN, HEATMAP_ZMAX + 0.001, 0.5)

    # Vertical: alto/positivo em cima, baixo/negativo embaixo.
    fig = plt.figure(figsize=(1.05, 4.2))
    cax = fig.add_axes([0.34, 0.08, 0.25, 0.84])
    cbar = fig.colorbar(sm, cax=cax, orientation="vertical", ticks=ticks)
    cbar.set_label("z-score", fontsize=9.0)
    cbar.ax.tick_params(labelsize=8.0)
    out_v = out_dir / f"{base}_zscore_SHARED_SCALE_VERTICAL.png"
    fig.savefig(out_v, dpi=PLOT_DPI)
    fig.savefig(out_v.with_suffix(".pdf"))
    plt.close(fig)

    # Horizontal: positivo à esquerda, como na prancha de referência.
    fig = plt.figure(figsize=(5.4, 1.05))
    cax = fig.add_axes([0.08, 0.30, 0.84, 0.26])
    cbar = fig.colorbar(sm, cax=cax, orientation="horizontal", ticks=ticks)
    cbar.ax.invert_xaxis()
    cbar.ax.xaxis.set_ticks_position("top")
    cbar.ax.xaxis.set_label_position("top")
    cbar.set_label("Z-SCORE", fontsize=9.5, labelpad=7)
    cbar.ax.tick_params(labelsize=8.0)
    out_h = out_dir / f"{base}_zscore_SHARED_SCALE_HORIZONTAL.png"
    fig.savefig(out_h, dpi=PLOT_DPI)
    fig.savefig(out_h.with_suffix(".pdf"))
    plt.close(fig)


def _format_supp_fdr(value) -> str:
    try:
        v = float(value)
    except Exception:
        return str(value)
    if not np.isfinite(v):
        return ""
    if v != 0 and abs(v) < 0.001:
        return f"{v:.3e}"
    return f"{v:.4f}"


def save_supplementary_category_tables(
    filt_sig: pd.DataFrame,
    selected_categories: list[str],
    out_dir: Path,
    base: str,
    fdr_cutoff: float,
    logfc_abs_cutoff: float,
) -> tuple[Path, Path]:
    """
    Cria uma tabela suplementar com TODOS os DEGs significativos das categorias
    selecionadas na GUI, usando a mesma classificação empregada nas figuras.

    Saídas:
      - .tabular único: Category, GeneID, Symbol, Description, logFC, FDR
      - .docx: uma subseção/tabela por categoria
    """
    cols = ["Category", "GeneID", "Symbol", "Description", "logFC", "FDR"]
    d = filt_sig[filt_sig["Category"].isin(selected_categories)].copy()
    for c in cols:
        if c not in d.columns:
            d[c] = ""

    # Preserva a ordem escolhida na GUI; dentro de cada categoria,
    # A-up (logFC positivo) aparece primeiro e B-up depois.
    d["Category"] = pd.Categorical(d["Category"], categories=selected_categories, ordered=True)
    d["logFC"] = pd.to_numeric(d["logFC"], errors="coerce")
    d["FDR"] = pd.to_numeric(d["FDR"], errors="coerce")
    d = d.sort_values(["Category", "logFC"], ascending=[True, False]).reset_index(drop=True)

    out_tab = out_dir / f"{base}_SUPPLEMENTARY_SelectedCategories_DEGs.tabular"
    d[cols].to_csv(out_tab, sep="\t", index=False, float_format="%.8g")

    try:
        from docx import Document
        from docx.shared import Mm, Pt
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError(
            "A tabela .tabular foi criada, mas para gerar a versão Word (.docx) "
            "é necessário instalar python-docx. No Windows, rode: "
            "py -m pip install python-docx"
        ) from e

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Mm(297)
    sec.page_height = Mm(210)
    sec.top_margin = Mm(10)
    sec.bottom_margin = Mm(10)
    sec.left_margin = Mm(10)
    sec.right_margin = Mm(10)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(14)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(11)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supplementary table — functionally classified significant DEGs")

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = info.add_run(
        f"Included genes satisfy FDR < {fdr_cutoff:g} and |logFC| ≥ {logfc_abs_cutoff:g}. "
        "Categories correspond to the predefined keyword-based classification used in the figures."
    )
    r.font.name = "Arial"
    r.font.size = Pt(8.5)

    present_categories = [c for c in selected_categories if c in set(d["Category"].astype(str))]
    for idx, cat in enumerate(present_categories):
        if idx > 0:
            doc.add_page_break()
        sub = d[d["Category"].astype(str) == cat].copy()
        h = doc.add_paragraph(style="Heading 2")
        h.add_run(f"{cat} (n = {len(sub)})")

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        widths = [Mm(27), Mm(28), Mm(155), Mm(25), Mm(25)]
        headers = ["GeneID", "Symbol", "Description", "logFC", "FDR"]
        hdr = table.rows[0]
        for i, (cell, header, width) in enumerate(zip(hdr.cells, headers, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p0 = cell.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = p0.add_run(header)
            rr.bold = True
            rr.font.name = "Arial"
            rr.font.size = Pt(8.5)

        # Repetir cabeçalho em páginas subsequentes.
        tr_pr = hdr._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

        for _, row in sub.iterrows():
            new_row = table.add_row()
            # Não permitir que uma linha da tabela seja quebrada entre duas páginas.
            row_pr = new_row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            row_pr.append(cant_split)
            cells = new_row.cells
            values = [
                str(row.get("GeneID", "")),
                str(row.get("Symbol", "")),
                str(row.get("Description", "")),
                "" if pd.isna(row.get("logFC")) else f"{float(row.get('logFC')):.3f}",
                _format_supp_fdr(row.get("FDR")),
            ]
            for j, (cell, value, width) in enumerate(zip(cells, values, widths)):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pp = cell.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 2 else WD_ALIGN_PARAGRAPH.CENTER
                rr = pp.add_run(value if value != "nan" else "")
                rr.font.name = "Arial"
                rr.font.size = Pt(8.0)

    out_docx = out_dir / f"{base}_SUPPLEMENTARY_SelectedCategories_DEGs.docx"
    doc.save(out_docx)
    return out_tab, out_docx


def save_corr_heatmap(corr: pd.DataFrame, out_png: Path, title: str):
    X = corr.values.astype(float)
    labels = corr.columns.tolist()

    fig, ax = plt.subplots(figsize=(CORR_WIDTH, CORR_HEIGHT))
    im = ax.imshow(X, aspect="auto", vmin=-1, vmax=1)
    cbar = fig.colorbar(im, ax=ax, fraction=0.05, pad=0.03)
    cbar.set_label("correlation", fontsize=8.5)
    cbar.ax.tick_params(labelsize=7.5)

    ax.set_xticks(np.arange(len(labels)))
    ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=7.5)
    ax.set_yticks(np.arange(len(labels)))
    ax.set_yticklabels(labels, fontsize=7.5)
    ax.set_title(title, fontsize=9.5, pad=7)

    fig.subplots_adjust(
        left=CORR_LEFT, right=CORR_RIGHT,
        bottom=CORR_BOTTOM, top=CORR_TOP
    )
    fig.savefig(out_png, dpi=PLOT_DPI)
    fig.savefig(out_png.with_suffix(".pdf"))
    plt.close(fig)


# ----------------------------
# BARPLOT DIVERGENTE POR CATEGORIA (UM ÚNICO)
# ----------------------------

def save_category_diverging_logfc_plot(
    df_cat: pd.DataFrame,
    out_png: Path,
    title: str,
    max_label_chars: int = MAX_LABEL_CHARS_COMPACT,
    label_mode: str = "desc",
):
    """
    Um único gráfico por categoria:
    - valor plotado = plotFC = -logFC (A-up fica negativo, B-up fica positivo)
    - y = descrição OU símbolo
    - ordenado pelo plotFC
    """
    if df_cat.empty:
        return

    d = df_cat.copy()
    d["logFC"] = pd.to_numeric(d["logFC"], errors="coerce")
    d = d.dropna(subset=["logFC"]).copy()
    if d.empty:
        return

    d["plotFC"] = -d["logFC"]  # A-up vira esquerda, B-up vira direita

    if label_mode == "symbol":
        labels = _build_symbol_labels(d["Symbol"], d.get("GeneID"), max_chars=max_label_chars)
    else:
        labels = _build_desc_labels(d["Description"], d.get("Symbol"), max_chars=max_label_chars)

    d["_label"] = labels
    d = d.sort_values("plotFC", ascending=True)

    n = len(d)

    # Largura fixa e estreita; altura cresce com o número de genes
    fig_h = BAR_TOP_IN + BAR_BOTTOM_IN + BAR_ROW_HEIGHT * n
    fig, ax = plt.subplots(figsize=(BARPLOT_WIDTH, fig_h))

    y = np.arange(n)
    ax.barh(y, d["plotFC"].astype(float).values, height=0.72)
    ax.set_yticks(y)
    ax.set_yticklabels(d["_label"].astype(str).values, fontsize=7.0)
    ax.axvline(0, linewidth=0.9)

    ax.set_xlabel("plotFC = -log2FC(A/B)  (A-up ← | → B-up)", fontsize=8.4)
    ax.set_title(title, fontsize=9.2, pad=7)
    ax.tick_params(axis="x", labelsize=7.8)
    ax.margins(y=0.01)

    fig.subplots_adjust(
        left=BAR_LEFT, right=BAR_RIGHT,
        bottom=BAR_BOTTOM_IN / fig_h,
        top=1.0 - (BAR_TOP_IN / fig_h)
    )
    fig.savefig(out_png, dpi=PLOT_DPI)
    fig.savefig(out_png.with_suffix(".pdf"))
    plt.close(fig)


def save_topN_up_plots_separate(
    filt_sig: pd.DataFrame,
    out_dir: Path,
    base: str,
    top_n: int,
    label_mode: str = "desc",
):
    """
    TopN A-up e TopN B-up em gráficos separados.
    label_mode: "desc" ou "symbol"
    """
    df = filt_sig.copy()
    df["logFC"] = pd.to_numeric(df["logFC"], errors="coerce")
    df = df.dropna(subset=["logFC"]).copy()

    suffix = "" if label_mode == "desc" else "_SYMBOL"

    # A-up
    dA = df[df["logFC"] > 0].sort_values("logFC", ascending=False).head(top_n).copy()
    if not dA.empty:
        if label_mode == "symbol":
            labA = _build_symbol_labels(dA["Symbol"], dA.get("GeneID"),
                                        max_chars=MAX_LABEL_CHARS_COMPACT)
            titleA = f"Top {top_n} A-up (SIG) — symbols"
        else:
            labA = _build_desc_labels(dA["Description"], dA.get("Symbol"),
                                      max_chars=MAX_LABEL_CHARS_COMPACT)
            titleA = f"Top {top_n} A-up (SIG) — descriptions"

        dA["_label"] = labA
        dA = dA.sort_values("logFC", ascending=True)
        nA = len(dA)

        fig_h = TOPN_TOP_IN + TOPN_BOTTOM_IN + TOPN_ROW_HEIGHT * nA
        fig, ax = plt.subplots(figsize=(TOPN_WIDTH, fig_h))
        y = np.arange(nA)
        ax.barh(y, dA["logFC"].astype(float).values, height=0.72)
        ax.set_yticks(y)
        ax.set_yticklabels(dA["_label"].values, fontsize=7.0)
        ax.set_xlabel("log2FC (A/B)", fontsize=8.4)
        ax.set_title(titleA, fontsize=9.2, pad=7)
        ax.tick_params(axis="x", labelsize=7.8)
        ax.margins(y=0.01)

        fig.subplots_adjust(
            left=BAR_LEFT, right=BAR_RIGHT,
            bottom=TOPN_BOTTOM_IN / fig_h,
            top=1.0 - (TOPN_TOP_IN / fig_h)
        )
        outA = out_dir / f"{base}_Top{top_n}_Aup{suffix}.png"
        fig.savefig(outA, dpi=PLOT_DPI)
        fig.savefig(outA.with_suffix(".pdf"))
        plt.close(fig)

    # B-up
    dB = df[df["logFC"] < 0].sort_values("logFC", ascending=True).head(top_n).copy()
    if not dB.empty:
        if label_mode == "symbol":
            labB = _build_symbol_labels(dB["Symbol"], dB.get("GeneID"),
                                        max_chars=MAX_LABEL_CHARS_COMPACT)
            titleB = f"Top {top_n} B-up (SIG) — symbols"
        else:
            labB = _build_desc_labels(dB["Description"], dB.get("Symbol"),
                                      max_chars=MAX_LABEL_CHARS_COMPACT)
            titleB = f"Top {top_n} B-up (SIG) — descriptions"

        dB["_label"] = labB
        dB["mag"] = (-dB["logFC"]).astype(float)
        dB = dB.sort_values("mag", ascending=True)
        nB = len(dB)

        fig_h = TOPN_TOP_IN + TOPN_BOTTOM_IN + TOPN_ROW_HEIGHT * nB
        fig, ax = plt.subplots(figsize=(TOPN_WIDTH, fig_h))
        y = np.arange(nB)
        ax.barh(y, dB["mag"].values, height=0.72)
        ax.set_yticks(y)
        ax.set_yticklabels(dB["_label"].values, fontsize=7.0)
        ax.set_xlabel("magnitude = -log2FC (A/B)  (B-up)", fontsize=8.4)
        ax.set_title(titleB, fontsize=9.2, pad=7)
        ax.tick_params(axis="x", labelsize=7.8)
        ax.margins(y=0.01)

        fig.subplots_adjust(
            left=BAR_LEFT, right=BAR_RIGHT,
            bottom=TOPN_BOTTOM_IN / fig_h,
            top=1.0 - (TOPN_TOP_IN / fig_h)
        )
        outB = out_dir / f"{base}_Top{top_n}_Bup{suffix}.png"
        fig.savefig(outB, dpi=PLOT_DPI)
        fig.savefig(outB.with_suffix(".pdf"))
        plt.close(fig)


# ----------------------------
# Pipeline principal
# ----------------------------

def _compute_profile_ylim(prof_df: pd.DataFrame, sample_cols: list[str]) -> tuple[float, float]:
    """Escala Y fixa para permitir comparação direta entre experimentos/fases."""
    return (PROFILE_YMIN, PROFILE_YMAX)


def run_all_analyses(
    edger_path: Path,
    counts_path: Path,
    out_dir: Path,
    base_name: str,
    fdr_cutoff: float = 0.05,
    logfc_abs_cutoff: float = 1.0,
    top_n_bar: int = 20,
    max_rows_heatmap: int = 300,
    drop_uncharacterized: bool = True,
    selected_profile_categories: list[str] | None = None,
):
    out_dir.mkdir(parents=True, exist_ok=True)
    base = sanitize_basename(base_name)
    if not base:
        raise ValueError("Nome base inválido.")

    # ---- Ler edgeR
    ed = _read_tabular(edger_path)
    geneid_col = _detect_column(ed, ["NCBI.GeneID", "GeneID", "geneid", "GeneId", "Gene ID"])
    sym_col    = _detect_column(ed, ["Symbol", "GeneName", "gene", "Name"])
    desc_col   = _detect_column(ed, ["Description", "descrição", "Descricao", "descricao", "Desc"])
    logfc_col  = _detect_column(ed, ["logFC", "log2FC", "log2FoldChange"])
    fdr_col    = _detect_column(ed, ["FDR", "adj.P.Val", "padj"])
    logcpm_col = _detect_column(ed, ["logCPM", "AveLogCPM", "logCPM."])

    missing = [("GeneID", geneid_col), ("logFC", logfc_col), ("FDR", fdr_col)]
    missing = [k for k, v in missing if v is None]
    if missing:
        raise ValueError(
            f"edgeR: faltando colunas obrigatórias: {missing}.\n"
            f"Colunas encontradas: {list(ed.columns)}"
        )

    rename_map = {
        geneid_col: "GeneID",
        logfc_col: "logFC",
        fdr_col: "FDR",
    }
    if sym_col:
        rename_map[sym_col] = "Symbol"
    if desc_col:
        rename_map[desc_col] = "Description"
    if logcpm_col:
        rename_map[logcpm_col] = "logCPM"

    ed = ed.rename(columns=rename_map)

    ed["GeneID"] = ed["GeneID"].astype(str)
    ed["logFC"] = _as_numeric_series(ed["logFC"])
    ed["FDR"] = _as_numeric_series(ed["FDR"])
    if "logCPM" in ed.columns:
        ed["logCPM"] = _as_numeric_series(ed["logCPM"])

    if "Symbol" not in ed.columns:
        ed["Symbol"] = ""
    if "Description" not in ed.columns:
        ed["Description"] = ""

    # ---- Ler counts
    ct = _read_tabular(counts_path)
    ct_geneid = _detect_column(ct, ["GeneID", "NCBI.GeneID", "geneid", "Gene ID"])
    if ct_geneid is None:
        ct_geneid = ct.columns[0]
    ct = ct.rename(columns={ct_geneid: "GeneID"})
    ct["GeneID"] = ct["GeneID"].astype(str)

    # Metadados opcionais vindos do counts (para fallback de símbolo/descrição)
    ct_sym_col = _detect_column(ct, ["Symbol", "symbol", "simbolo", "símbolo", "GeneName", "gene", "Name"])
    ct_desc_col = _detect_column(ct, ["Description", "descrição", "Descricao", "descricao", "Desc"])
    meta_cols = ["GeneID"]
    rename_meta = {}
    if ct_sym_col and ct_sym_col != "GeneID":
        meta_cols.append(ct_sym_col)
        rename_meta[ct_sym_col] = "Symbol_ct"
    if ct_desc_col and ct_desc_col != "GeneID":
        meta_cols.append(ct_desc_col)
        rename_meta[ct_desc_col] = "Description_ct"

    ct_meta = ct[meta_cols].copy().rename(columns=rename_meta)

    # Detectar e normalizar A1..B3/4/5
    ct, sample_cols = detect_and_normalize_sample_cols(ct)

    # manter GeneID + amostras, mas preservar metadados do counts
    ct = ct[["GeneID"] + sample_cols].copy()
    ct[sample_cols] = ct[sample_cols].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    ct = ct.merge(ct_meta, on="GeneID", how="left")

    # ---- Merge edgeR + counts
    m = ed.merge(ct, on="GeneID", how="left")
    for c in sample_cols:
        if c not in m.columns:
            m[c] = 0.0
        m[c] = pd.to_numeric(m[c], errors="coerce").fillna(0.0)

    # ---- Preencher Symbol/Description ausentes com dados do counts (se houver)
    if "Symbol_ct" in m.columns:
        sym_main = m["Symbol"].fillna("").astype(str).replace("nan", "").str.strip()
        sym_ct = m["Symbol_ct"].fillna("").astype(str).replace("nan", "").str.strip()
        empty = sym_main.eq("") | sym_main.str.lower().eq("nan")
        m.loc[empty, "Symbol"] = sym_ct.loc[empty]

    if "Description_ct" in m.columns:
        desc_main = m["Description"].fillna("").astype(str).replace("nan", "").str.strip()
        desc_ct = m["Description_ct"].fillna("").astype(str).replace("nan", "").str.strip()
        empty = desc_main.eq("") | desc_main.str.lower().eq("nan")
        m.loc[empty, "Description"] = desc_ct.loc[empty]

    # ---- logCPM (se edgeR não tiver)
    if "logCPM" not in m.columns:
        log2cpm = compute_log2cpm_from_counts(m[["GeneID"] + sample_cols], sample_cols)
        tmp = m.merge(log2cpm, on="GeneID", how="left", suffixes=("", "_log2cpm"))
        m["logCPM"] = tmp[sample_cols].mean(axis=1)

    # ---- Direção
    m["Direction"] = np.where(m["logFC"] > 0, "A-up", np.where(m["logFC"] < 0, "B-up", "0"))

    # ---- Categorias
    m["CategoryList"] = [
        categorize_row(gid, sym, desc)
        for gid, sym, desc in zip(
            m["GeneID"].fillna("").astype(str),
            m["Symbol"].fillna("").astype(str),
            m["Description"].fillna("").astype(str),
        )
    ]
    m["Category"] = m["CategoryList"].apply(lambda x: x[0] if isinstance(x, list) and x else "Other")

    # ---- Filtro estrito SIG
    filt_sig = m[(m["FDR"] < fdr_cutoff) & (m["logFC"].abs() >= logfc_abs_cutoff)].copy()

    # ---- Excluir descrições genéricas
    if drop_uncharacterized:
        filt_sig = filter_informative_only(filt_sig)

    # ---- Saídas tabulares (SIG)
    out_all = out_dir / f"{base}_SIG_FDR{fdr_cutoff:g}_logFC{logfc_abs_cutoff:g}_ALL.tabular"
    out_A   = out_dir / f"{base}_SIG_FDR{fdr_cutoff:g}_logFC{logfc_abs_cutoff:g}_Aup.tabular"
    out_B   = out_dir / f"{base}_SIG_FDR{fdr_cutoff:g}_logFC{logfc_abs_cutoff:g}_Bup.tabular"

    head_cols = ["GeneID", "Symbol", "Description", "logFC", "FDR", "logCPM", "Direction", "Category"]
    rest_cols = [c for c in filt_sig.columns if c not in head_cols]
    filt_sig = filt_sig[head_cols + rest_cols]

    filt_sig.to_csv(out_all, sep="\t", index=False)
    filt_sig[filt_sig["Direction"] == "A-up"].to_csv(out_A, sep="\t", index=False)
    filt_sig[filt_sig["Direction"] == "B-up"].to_csv(out_B, sep="\t", index=False)

    summary = (
        filt_sig.groupby(["Category", "Direction"], as_index=False)
               .size()
               .pivot(index="Category", columns="Direction", values="size")
               .fillna(0)
               .astype(int)
               .reset_index()
    )
    out_sum = out_dir / f"{base}_summary_category_by_direction.tabular"
    summary.to_csv(out_sum, sep="\t", index=False)

    # ---- TOP A-up e TOP B-up (descrição e símbolo)
    save_topN_up_plots_separate(
        filt_sig=filt_sig,
        out_dir=out_dir,
        base=base,
        top_n=top_n_bar,
        label_mode="desc",
    )
    save_topN_up_plots_separate(
        filt_sig=filt_sig,
        out_dir=out_dir,
        base=base,
        top_n=top_n_bar,
        label_mode="symbol",
    )

    # ---- Categorias selecionadas na GUI
    # A mesma seleção controla:
    #   - barplots por categoria
    #   - heatmaps por categoria
    #   - painéis combinados barplot + heatmap
    #   - gráfico final de perfis e correlação
    # As saídas gerais (tabelas, TopN e heatmap SIG geral) continuam completas.
    if selected_profile_categories is None:
        selected_profile_categories = list(CATEGORY_NAMES)
    selected_profile_categories = [c for c in selected_profile_categories if c in SELECTABLE_CATEGORIES]
    selected_category_set = set(selected_profile_categories)

    # ---- Escalas z-score compartilhadas para montagem das pranchas
    save_shared_zscore_scales(out_dir=out_dir, base=base)

    # ---- Tabela suplementar das categorias selecionadas
    out_supp_tab, out_supp_docx = save_supplementary_category_tables(
        filt_sig=filt_sig,
        selected_categories=selected_profile_categories,
        out_dir=out_dir,
        base=base,
        fdr_cutoff=fdr_cutoff,
        logfc_abs_cutoff=logfc_abs_cutoff,
    )

    # ---- Gráficos por categoria (descrição e símbolo)
    cats_present_sig = set(filt_sig["Category"].dropna().unique().tolist())
    cats = [cat for cat in SELECTABLE_CATEGORIES if cat in cats_present_sig and cat in selected_category_set]
    for cat in cats:
        sub = filt_sig[filt_sig["Category"] == cat].copy()
        if len(sub) < 1:
            continue

        out_png_desc = out_dir / f"{base}_Category_{cat}_logFC_diverging.png"
        save_category_diverging_logfc_plot(
            df_cat=sub,
            out_png=out_png_desc,
            title=f"{cat} (SIG) — plotFC=-logFC (A-up esquerda; B-up direita) — n={len(sub)} — descriptions",
            max_label_chars=MAX_LABEL_CHARS_COMPACT,
            label_mode="desc",
        )

        out_png_sym = out_dir / f"{base}_Category_{cat}_logFC_diverging_SYMBOL.png"
        save_category_diverging_logfc_plot(
            df_cat=sub,
            out_png=out_png_sym,
            title=f"{cat} (SIG) — plotFC=-logFC (A-up esquerda; B-up direita) — n={len(sub)} — symbols",
            max_label_chars=MAX_LABEL_CHARS_COMPACT,
            label_mode="symbol",
        )

    # ---- Heatmaps
    log2cpm_mat = compute_log2cpm_from_counts(m[["GeneID"] + sample_cols], sample_cols)
    hm = filt_sig[["GeneID", "Symbol", "Description", "logFC", "Category", "Direction"]].merge(
        log2cpm_mat, on="GeneID", how="left"
    ).fillna(0.0)

    hm_z = zscore_rows(hm, sample_cols)

    # Ordenar pelo logFC para manter a MESMA ordem visual dos barplots:
    # no barplot, plotFC = -logFC e os maiores plotFC aparecem no topo.
    # Portanto, no heatmap (que desenha a primeira linha no topo),
    # ordenar logFC do menor para o maior reproduz a mesma ordem:
    # B-up no topo -> A-up embaixo.
    hm_z = hm_z.sort_values("logFC", ascending=True).reset_index(drop=True)

    # Heatmap SIG geral - descrição
    rowlab_sig_desc = _build_desc_labels(hm_z["Description"], hm_z["Symbol"], max_chars=MAX_LABEL_CHARS_COMPACT)
    save_heatmap(
        hm_z,
        out_png=out_dir / f"{base}_heatmap_SIG_zscore.png",
        title=f"Heatmap (z-score por gene) — SIG: FDR<{fdr_cutoff:g} & |logFC|>={logfc_abs_cutoff:g} — descriptions",
        sample_cols=sample_cols,
        row_labels=rowlab_sig_desc,
        max_rows=max_rows_heatmap,
    )

    # Heatmap SIG geral - símbolo
    rowlab_sig_sym = _build_symbol_labels(hm_z["Symbol"], hm_z["GeneID"], max_chars=MAX_LABEL_CHARS_COMPACT)
    save_heatmap(
        hm_z,
        out_png=out_dir / f"{base}_heatmap_SIG_zscore_SYMBOL.png",
        title=f"Heatmap (z-score por gene) — SIG: FDR<{fdr_cutoff:g} & |logFC|>={logfc_abs_cutoff:g} — symbols",
        sample_cols=sample_cols,
        row_labels=rowlab_sig_sym,
        max_rows=max_rows_heatmap,
    )

    # Heatmaps por categoria (SIG) - descrição e símbolo
    cats_present_hm = set(hm_z["Category"].unique())
    hm_categories = [cat for cat in SELECTABLE_CATEGORIES if cat in cats_present_hm and cat in selected_category_set]
    for cat in hm_categories:
        sub = hm_z[hm_z["Category"] == cat].copy()
        if len(sub) < 1:
            continue
        sub = sub.sort_values("logFC", ascending=True).reset_index(drop=True)

        rowlab_cat_desc = _build_desc_labels(sub["Description"], sub["Symbol"], max_chars=MAX_LABEL_CHARS_COMPACT)
        save_heatmap(
            sub,
            out_png=out_dir / f"{base}_heatmap_{cat}_zscore.png",
            title=f"Heatmap (z-score) — {cat} (n={len(sub)}) — descriptions",
            sample_cols=sample_cols,
            row_labels=rowlab_cat_desc,
            max_rows=max_rows_heatmap,
        )

        rowlab_cat_sym = _build_symbol_labels(sub["Symbol"], sub["GeneID"], max_chars=MAX_LABEL_CHARS_COMPACT)
        save_heatmap(
            sub,
            out_png=out_dir / f"{base}_heatmap_{cat}_zscore_SYMBOL.png",
            title=f"Heatmap (z-score) — {cat} (n={len(sub)}) — symbols",
            sample_cols=sample_cols,
            row_labels=rowlab_cat_sym,
            max_rows=max_rows_heatmap,
        )

        # Painel combinado: barplot + heatmap, gene a gene
        save_combined_bar_heatmap(
            sub,
            out_png=out_dir / f"{base}_Category_{cat}_COMBINED.png",
            title=f"{cat} (SIG) — n={len(sub)}",
            sample_cols=sample_cols,
            label_mode="desc",
            max_label_chars=MAX_LABEL_CHARS_COMPACT,
        )

        save_combined_bar_heatmap(
            sub,
            out_png=out_dir / f"{base}_Category_{cat}_COMBINED_SYMBOL.png",
            title=f"{cat} (SIG) — n={len(sub)}",
            sample_cols=sample_cols,
            label_mode="symbol",
            max_label_chars=MAX_LABEL_CHARS_COMPACT,
        )

    # ---- Perfis por categoria (SIG) e correlação
    # Usa a MESMA seleção da GUI aplicada aos gráficos por categoria.
    profiles = []

    cats_present = set(hm_z["Category"].unique())
    ordered_present_cats = [c for c in SELECTABLE_CATEGORIES if c in cats_present]
    for cat in ordered_present_cats:
        sub = hm_z[hm_z["Category"] == cat]
        if len(sub) < 1:
            continue
        mean_profile = sub[sample_cols].astype(float).mean(axis=0)
        row = {"Category": cat, **{k: float(v) for k, v in mean_profile.items()}, "n_genes": int(len(sub))}
        profiles.append(row)

    prof_df = pd.DataFrame(profiles)
    out_prof = out_dir / f"{base}_category_profiles_zscore_SIG.tabular"
    if not prof_df.empty:
        prof_df.to_csv(out_prof, sep="	", index=False)

        prof_plot_df = prof_df[prof_df["Category"].isin(selected_profile_categories)].copy()
        prof_plot_df["Category"] = pd.Categorical(prof_plot_df["Category"], categories=selected_profile_categories, ordered=True)
        prof_plot_df = prof_plot_df.sort_values("Category").reset_index(drop=True)

        if not prof_plot_df.empty:
            corr = prof_plot_df.set_index("Category")[sample_cols].T.corr()
            out_corr = out_dir / f"{base}_category_profile_correlation.tabular"
            corr.to_csv(out_corr, sep="	")
            save_corr_heatmap(
                corr,
                out_png=out_dir / f"{base}_category_profile_correlation.png",
                title="Correlação entre categorias (perfil médio z-score, genes SIG)",
            )

            y0, y1 = _compute_profile_ylim(prof_plot_df, sample_cols)

            def _draw_profile_lines(ax, include_legend: bool):
                for _, r in prof_plot_df.iterrows():
                    cat = str(r["Category"])
                    color = PROFILE_CATEGORY_COLORS.get(cat, None)
                    short = PROFILE_CATEGORY_LABELS.get(cat, cat)
                    ax.plot(
                        sample_cols,
                        [r[c] for c in sample_cols],
                        marker="o",
                        markersize=PROFILE_MARKER_SIZE,
                        linewidth=1.8,
                        color=color,
                        label=f"{short} (n={r['n_genes']})",
                    )
                ax.axhline(0, color="0.55", linewidth=0.8)
                ax.set_xticks(range(len(sample_cols)))
                ax.set_xticklabels(sample_cols, rotation=45, ha="right")
                ax.set_ylim(y0, y1)
                ax.set_ylabel("mean z-score")
                ax.set_title("Perfis médios por categoria (genes SIG)")

                if include_legend:
                    legend_n = len(prof_plot_df)
                    legend_font = 7.0 if legend_n <= 6 else 6.5
                    ax.legend(
                        fontsize=legend_font,
                        loc="upper left",
                        bbox_to_anchor=(1.01, 1.0),
                        borderaxespad=0.0,
                        frameon=True,
                    )

            # ----------------------------------------------------
            # 1) Versão standalone: legenda externa em canvas FIXO
            # NÃO usar bbox_inches='tight': ele altera a largura final
            # conforme o conteúdo/comprimento da legenda.
            # ----------------------------------------------------
            fig, ax = plt.subplots(figsize=(PROFILE_STANDALONE_WIDTH, PROFILE_HEIGHT))
            _draw_profile_lines(ax, include_legend=True)
            fig.subplots_adjust(
                left=PROFILE_LEFT,
                right=PROFILE_STANDALONE_RIGHT,
                bottom=PROFILE_BOTTOM,
                top=PROFILE_TOP,
            )
            out_lines = out_dir / f"{base}_category_profiles_lines.png"
            fig.savefig(out_lines, dpi=PLOT_DPI)
            fig.savefig(out_lines.with_suffix(".pdf"))
            plt.close(fig)

            # ----------------------------------------------------
            # 2) Versão PANEL: sem legenda, mesmo tamanho/área útil
            # em TODAS as comparações. Esta é a versão recomendada
            # para montar pranchas 2x2 no artigo.
            # ----------------------------------------------------
            fig, ax = plt.subplots(figsize=(PROFILE_PANEL_WIDTH, PROFILE_HEIGHT))
            _draw_profile_lines(ax, include_legend=False)
            fig.subplots_adjust(
                left=PROFILE_LEFT,
                right=PROFILE_PANEL_RIGHT,
                bottom=PROFILE_BOTTOM,
                top=PROFILE_TOP,
            )
            out_panel = out_dir / f"{base}_category_profiles_lines_PANEL.png"
            fig.savefig(out_panel, dpi=PLOT_DPI)
            fig.savefig(out_panel.with_suffix(".pdf"))
            plt.close(fig)

            # ----------------------------------------------------
            # 3) Legenda separada/compartilhada para a prancha.
            # Sem n, porque n varia entre as comparações/painéis.
            # Use apenas UMA vez na figura 2x2.
            # ----------------------------------------------------
            legend_cats = [
                c for c in selected_profile_categories
                if c in PROFILE_CATEGORY_COLORS
            ]
            handles = [
                Line2D(
                    [0], [0],
                    color=PROFILE_CATEGORY_COLORS[c],
                    marker="o",
                    linewidth=1.8,
                    markersize=PROFILE_MARKER_SIZE,
                    label=PROFILE_CATEGORY_LABELS.get(c, c),
                )
                for c in legend_cats
            ]
            if handles:
                legend_w = 7.2
                legend_h = 0.65 if len(handles) <= 5 else 0.95
                fig_leg = plt.figure(figsize=(legend_w, legend_h))
                fig_leg.legend(
                    handles=handles,
                    loc="center",
                    ncol=min(5, len(handles)),
                    frameon=False,
                    fontsize=8.0,
                )
                out_leg = out_dir / f"{base}_category_profiles_SHARED_LEGEND.png"
                fig_leg.savefig(out_leg, dpi=PLOT_DPI)
                fig_leg.savefig(out_leg.with_suffix(".pdf"))
                plt.close(fig_leg)


    return {
        "n_total_edger": int(len(m)),
        "n_sig_after_filters": int(len(filt_sig)),
        "drop_uncharacterized": bool(drop_uncharacterized),
        "selected_profile_categories": list(selected_profile_categories),
        "out_dir": str(out_dir),
        "base": base,
        "sample_cols": sample_cols,
        "supplementary_tabular": str(out_supp_tab),
        "supplementary_docx": str(out_supp_docx),
    }



# ----------------------------
# GUI
# ----------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("edgeR Analyzer — A vs B (GUI)")

        self.edger_var = tk.StringVar()
        self.counts_var = tk.StringVar()
        self.outdir_var = tk.StringVar()
        self.base_var = tk.StringVar(value="A_vs_B")

        self.fdr_var = tk.StringVar(value="0.05")
        self.logfc_var = tk.StringVar(value="1.0")
        self.topn_var = tk.StringVar(value="20")
        self.maxrows_var = tk.StringVar(value="300")

        self.drop_unchar_var = tk.BooleanVar(value=True)
        # Por padrão, gera todas as categorias funcionais e deixa "Other" desmarcada,
        # pois ela costuma conter muitos genes e produzir figuras muito grandes.
        self.category_vars = {
            cat: tk.BooleanVar(value=(cat != "Other"))
            for cat in SELECTABLE_CATEGORIES
        }

        pad = {"padx": 8, "pady": 6}

        frm = ttk.Frame(self)
        frm.pack(fill="both", expand=True)

        ttk.Label(frm, text="edgeR result (.tabular):").grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.edger_var, width=78).grid(row=0, column=1, **pad)
        ttk.Button(frm, text="Selecionar…", command=self.pick_edger).grid(row=0, column=2, **pad)

        ttk.Label(frm, text="Counts (GeneID + A1..B3/4/5) ou (GeneID + ... + Counts_A1..Counts_B3/4/5):").grid(row=1, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.counts_var, width=78).grid(row=1, column=1, **pad)
        ttk.Button(frm, text="Selecionar…", command=self.pick_counts).grid(row=1, column=2, **pad)

        ttk.Label(frm, text="Pasta de saída:").grid(row=2, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.outdir_var, width=78).grid(row=2, column=1, **pad)
        ttk.Button(frm, text="Selecionar…", command=self.pick_outdir).grid(row=2, column=2, **pad)

        ttk.Label(frm, text="Nome base do experimento:").grid(row=3, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.base_var, width=78).grid(row=3, column=1, **pad)

        pfrm = ttk.LabelFrame(frm, text="Parâmetros (padrão: estrito)")
        pfrm.grid(row=4, column=0, columnspan=3, sticky="we", **pad)

        ttk.Label(pfrm, text="FDR cutoff:").grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(pfrm, textvariable=self.fdr_var, width=10).grid(row=0, column=1, sticky="w", **pad)

        ttk.Label(pfrm, text="|logFC| cutoff:").grid(row=0, column=2, sticky="w", **pad)
        ttk.Entry(pfrm, textvariable=self.logfc_var, width=10).grid(row=0, column=3, sticky="w", **pad)

        ttk.Label(pfrm, text="Top N (Top A-up e Top B-up):").grid(row=1, column=0, sticky="w", **pad)
        ttk.Entry(pfrm, textvariable=self.topn_var, width=10).grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(pfrm, text="Max genes no heatmap:").grid(row=1, column=2, sticky="w", **pad)
        ttk.Entry(pfrm, textvariable=self.maxrows_var, width=10).grid(row=1, column=3, sticky="w", **pad)

        ttk.Checkbutton(
            pfrm,
            text="Excluir descrições genéricas (uncharacterized/hypothetical/predicted/unknown...)",
            variable=self.drop_unchar_var
        ).grid(row=2, column=0, columnspan=4, sticky="w", **pad)

        cfrm = ttk.LabelFrame(frm, text="Categorias a gerar nos gráficos por categoria")
        cfrm.grid(row=5, column=0, columnspan=3, sticky="we", **pad)

        ttk.Label(
            cfrm,
            text=(
                "A seleção vale para barplots, heatmaps, imagens combinadas (descrição e símbolo), "
                "perfil final, correlação e tabela suplementar. TopN e o heatmap SIG geral permanecem completos."
            ),
        ).grid(row=0, column=0, columnspan=3, sticky="w", padx=8, pady=(6, 2))

        for i, cat in enumerate(SELECTABLE_CATEGORIES):
            ttk.Checkbutton(cfrm, text=cat, variable=self.category_vars[cat]).grid(
                row=1 + (i // 3), column=i % 3, sticky="w", padx=8, pady=2
            )

        button_row = 1 + ((len(SELECTABLE_CATEGORIES) - 1) // 3) + 1
        ttk.Button(cfrm, text="Padrão (sem Other)", command=self.reset_category_selection).grid(
            row=button_row, column=0, sticky="w", padx=8, pady=(6, 8)
        )
        ttk.Button(cfrm, text="Selecionar todas", command=self.select_all_categories).grid(
            row=button_row, column=1, sticky="w", padx=8, pady=(6, 8)
        )
        ttk.Button(cfrm, text="Limpar seleção", command=self.clear_category_selection).grid(
            row=button_row, column=2, sticky="w", padx=8, pady=(6, 8)
        )

        bfrm = ttk.Frame(frm)
        bfrm.grid(row=6, column=0, columnspan=3, sticky="e", **pad)

        ttk.Button(bfrm, text="Rodar análises", command=self.run).grid(row=0, column=0, **pad)
        ttk.Button(bfrm, text="Clear", command=self.clear).grid(row=0, column=1, **pad)

        self.status = tk.StringVar(value="Pronto.")
        ttk.Label(frm, textvariable=self.status).grid(row=7, column=0, columnspan=3, sticky="w", **pad)

    def reset_category_selection(self):
        """Padrão: todas as categorias funcionais marcadas e Other desmarcada."""
        for cat, var in self.category_vars.items():
            var.set(cat != "Other")

    def select_all_categories(self):
        for var in self.category_vars.values():
            var.set(True)

    def clear_category_selection(self):
        for var in self.category_vars.values():
            var.set(False)

    def pick_edger(self):
        p = filedialog.askopenfilename(
            title="Selecione o arquivo edgeR (.tabular)",
            filetypes=[("Tabular/TSV", "*.tabular *.tsv *.txt"), ("All files", "*.*")]
        )
        if p:
            self.edger_var.set(p)
            if not self.base_var.get().strip():
                self.base_var.set(Path(p).stem)

    def pick_counts(self):
        p = filedialog.askopenfilename(
            title="Selecione o arquivo counts (edgeReady)",
            filetypes=[("Tabular/TSV", "*.tabular *.tsv *.txt"), ("All files", "*.*")]
        )
        if p:
            self.counts_var.set(p)

    def pick_outdir(self):
        p = filedialog.askdirectory(title="Selecione a pasta de saída")
        if p:
            self.outdir_var.set(p)

    def clear(self):
        self.edger_var.set("")
        self.counts_var.set("")
        self.outdir_var.set("")
        self.base_var.set("A_vs_B")
        self.fdr_var.set("0.05")
        self.logfc_var.set("1.0")
        self.topn_var.set("20")
        self.maxrows_var.set("300")
        self.drop_unchar_var.set(True)
        self.reset_category_selection()
        self.status.set("Pronto.")

    def run(self):
        try:
            edger = Path(self.edger_var.get().strip()).expanduser()
            counts = Path(self.counts_var.get().strip()).expanduser()
            outdir_str = self.outdir_var.get().strip()
            base = self.base_var.get().strip()

            if not edger.exists():
                raise FileNotFoundError("edgeR: arquivo não encontrado.")
            if not counts.exists():
                raise FileNotFoundError("Counts: arquivo não encontrado.")
            if not outdir_str:
                raise ValueError("Selecione a pasta de saída.")
            outdir = Path(outdir_str).expanduser()

            fdr = float(self.fdr_var.get().strip().replace(",", "."))
            logfc = float(self.logfc_var.get().strip().replace(",", "."))
            topn = int(float(self.topn_var.get().strip().replace(",", ".")))
            maxrows = int(float(self.maxrows_var.get().strip().replace(",", ".")))

            if fdr <= 0 or fdr >= 1:
                raise ValueError("FDR cutoff deve ser entre 0 e 1.")
            if logfc < 0:
                raise ValueError("|logFC| cutoff deve ser >= 0.")
            if topn < 1:
                raise ValueError("Top N deve ser >= 1.")
            if maxrows < 10:
                raise ValueError("Max genes no heatmap deve ser >= 10.")

            selected_profile_categories = [
                cat for cat in SELECTABLE_CATEGORIES
                if self.category_vars[cat].get()
            ]
            if not selected_profile_categories:
                raise ValueError("Selecione pelo menos uma categoria para gerar os gráficos por categoria/perfis.")

            self.status.set("Rodando…")
            self.update_idletasks()

            res = run_all_analyses(
                edger_path=edger,
                counts_path=counts,
                out_dir=outdir,
                base_name=base,
                fdr_cutoff=fdr,
                logfc_abs_cutoff=logfc,
                top_n_bar=topn,
                max_rows_heatmap=maxrows,
                drop_uncharacterized=self.drop_unchar_var.get(),
                selected_profile_categories=selected_profile_categories,
            )

            self.status.set("Concluído.")
            msg_extra = "SIM" if res["drop_uncharacterized"] else "NÃO"
            messagebox.showinfo(
                "OK",
                "Concluído!\n\n"
                f"Genes edgeR: {res['n_total_edger']}\n"
                f"Genes SIG (após filtros): {res['n_sig_after_filters']}\n"
                f"Excluir descrições genéricas: {msg_extra}\n"
                f"Amostras detectadas: {', '.join(res['sample_cols'])}\n"
                f"Categorias geradas: {', '.join(res['selected_profile_categories'])}\n\n"
                f"Pasta:\n{res['out_dir']}\n\n"
                "Saídas principais:\n"
                f"- {res['base']}_Top{topn}_Aup.png\n"
                f"- {res['base']}_Top{topn}_Bup.png\n"
                f"- {res['base']}_Category_<Categoria>_logFC_diverging.png\n"
                f"- {res['base']}_heatmap_SIG_zscore.png\n"
                f"- {res['base']}_heatmap_<Categoria>_zscore.png\n"
                f"- {res['base']}_Category_<Categoria>_COMBINED.png\n"
                f"- {res['base']}_Category_<Categoria>_COMBINED_SYMBOL.png\n"
                f"- {res['base']}_zscore_SHARED_SCALE_HORIZONTAL.png\n"
                f"- {res['base']}_zscore_SHARED_SCALE_VERTICAL.png\n"
                f"- {res['base']}_SUPPLEMENTARY_SelectedCategories_DEGs.tabular\n"
                f"- {res['base']}_SUPPLEMENTARY_SelectedCategories_DEGs.docx\n"
            )

        except Exception as e:
            self.status.set("Erro.")
            messagebox.showerror("Erro", str(e))


if __name__ == "__main__":
    App().mainloop()